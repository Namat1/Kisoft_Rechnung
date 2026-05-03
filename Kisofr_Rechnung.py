import io
import math
import re
from datetime import datetime
from typing import Dict, List, Tuple

import pandas as pd
import streamlit as st
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


# -----------------------------------------------------------------------------
# Grundeinstellung
# -----------------------------------------------------------------------------
st.set_page_config(
    page_title="Auftragspool Auswertung",
    page_icon="📦",
    layout="wide",
)

st.markdown(
    """
    <style>
        .block-container {
            padding-top: 1.2rem;
            padding-bottom: 1.2rem;
            max-width: 1500px;
        }
        h1, h2, h3 {
            letter-spacing: -0.03em;
        }
        div[data-testid="stMetric"] {
            background: #f7f7f8;
            border: 1px solid #e5e7eb;
            padding: 14px 16px;
            border-radius: 14px;
        }
        .small-note {
            color: #6b7280;
            font-size: 0.92rem;
        }
    </style>
    """,
    unsafe_allow_html=True,
)


# -----------------------------------------------------------------------------
# Spaltenlogik
# -----------------------------------------------------------------------------
OUTPUT_COLUMNS = [
    "CSB Tournummer",
    "Einheit",
    "CSB Kundennummer",
    "Kundenname",
    "Stadt",
    "Straße",
    "Anzahl gepl. E2",
    "Anzahl gepl. E1",
    "Anzahl gepl. KARTON",
    "Anzahl Rolli/TKT",
]

NUMBER_COLUMNS = [
    "Anzahl gepl. E2",
    "Anzahl gepl. E1",
    "Anzahl gepl. KARTON",
]

CALCULATED_COLUMN = "Anzahl Rolli/TKT"

# TKT-Touren: Diese Nummern und Nummern mit gleichem Aufbau werden als TKT gerechnet.
# Beispiel gleicher Aufbau: 12221, 22221, 32221 ... beziehungsweise 17779, 27779 ...
TKT_EXACT_TOURS = {
    "12221",
    "12222",
    "12223",
    "14444",
    "17773",
    "17778",
    "17779",
    "27779",
}

TKT_SUFFIXES = {
    "2221",
    "2222",
    "2223",
    "4444",
    "7773",
    "7778",
    "7779",
}

COLUMN_ALIASES: Dict[str, List[str]] = {
    "CSB Tournummer": ["CSB Tournummer", "CSB-Tournummer", "CSB Tour", "Tournummer CSB"],
    "CSB Kundennummer": ["CSB Kundennummer", "CSB-Kundennummer", "CSB Kunden Nummer", "CSB Kunden-Nr", "CSB Kundennr"],
    "Kundenname": ["Kundenname", "Kunden Name", "Name", "Kunde"],
    "Stadt": ["Stadt", "Ort"],
    "Straße": ["Straße", "Strasse", "Str.", "Straße / Hausnummer", "Strasse / Hausnummer"],
    "Anzahl gepl. E2": ["Anzahl gepl. E2", "Anzahl E2", "E2"],
    "Anzahl gepl. E1": ["Anzahl gepl. E1", "Anzahl E1", "E1"],
    "Anzahl gepl. KARTON": ["Anzahl gepl. KARTON", "Anzahl gepl. KT", "Anzahl KARTON", "Anzahl KT", "KARTON", "KT"],
}


def normalize_column_name(value: object) -> str:
    """Macht Spaltennamen vergleichbar, ohne die Originalspalten zu verändern."""
    text = str(value).replace("\ufeff", "").strip()
    text = re.sub(r"\.\d+$", "", text)  # pandas-Zusatz bei doppelten Spaltennamen entfernen
    text = text.replace("ß", "ss")
    text = re.sub(r"\s+", " ", text)
    return text.casefold()


def find_required_columns(df: pd.DataFrame) -> Tuple[Dict[str, str], List[str]]:
    normalized_to_original: Dict[str, str] = {}
    for original in df.columns:
        normalized = normalize_column_name(original)
        normalized_to_original.setdefault(normalized, original)

    found: Dict[str, str] = {}
    missing: List[str] = []

    for target, aliases in COLUMN_ALIASES.items():
        match = None
        for alias in aliases:
            normalized_alias = normalize_column_name(alias)
            if normalized_alias in normalized_to_original:
                match = normalized_to_original[normalized_alias]
                break
        if match is None:
            missing.append(target)
        else:
            found[target] = match

    return found, missing


# -----------------------------------------------------------------------------
# Lesen und Berechnen
# -----------------------------------------------------------------------------
def read_csv_robust(uploaded_file) -> pd.DataFrame:
    """Liest den Kisoft Auftragspool robust ein."""
    raw = uploaded_file.getvalue()
    encodings = ["utf-8-sig", "utf-8", "cp1252", "latin1"]
    last_error = None

    for encoding in encodings:
        try:
            return pd.read_csv(
                io.BytesIO(raw),
                sep=";",
                encoding=encoding,
                dtype=str,
                keep_default_na=False,
            )
        except Exception as error:
            last_error = error

    raise ValueError(f"Die Datei konnte nicht gelesen werden: {last_error}")


def clean_text(value: object) -> str:
    if pd.isna(value):
        return ""
    text = str(value).strip()
    if text.lower() in {"nan", "none", "nat"}:
        return ""
    return text


def clean_customer_number(value: object) -> str:
    """Macht aus 12345.0 wieder 12345, lässt echte führende Nullen aber unverändert."""
    text = clean_text(value)
    if text.endswith(".0") and text[:-2].isdigit():
        return text[:-2]
    return text


def normalize_tour_number(value: object) -> str:
    """Bereitet die Tournummer für Regeln vor."""
    text = clean_customer_number(value)
    return re.sub(r"\D", "", text)


def is_tkt_tour(value: object) -> bool:
    """
    Erkennt TKT-Touren.

    Exakte Touren werden erkannt.
    Zusätzlich wird der gleiche Aufbau erkannt: fünf Ziffern, erste Ziffer variabel,
    die letzten vier Ziffern entsprechen einem bekannten TKT-Muster.
    """
    digits = normalize_tour_number(value)
    if digits in TKT_EXACT_TOURS:
        return True
    return len(digits) == 5 and digits[1:] in TKT_SUFFIXES


def parse_number(value: object) -> float:
    """Wandelt deutsche und internationale Zahlenformate in Zahlen um."""
    text = clean_text(value)
    if not text:
        return 0.0

    text = text.replace(" ", "")
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".")
    elif "," in text:
        text = text.replace(",", ".")

    try:
        return float(text)
    except ValueError:
        return 0.0


def round_up_full(value: object) -> int:
    """Rundet auf volle Einheiten auf. Null bleibt Null."""
    try:
        number = float(value)
    except (TypeError, ValueError):
        return 0
    if number <= 0:
        return 0
    return int(math.ceil(number))


def sort_key(series: pd.Series) -> pd.Series:
    numeric = pd.to_numeric(series, errors="coerce")
    return numeric.fillna(999999999).astype(float)


def sort_result(df: pd.DataFrame) -> pd.DataFrame:
    result = df.copy()
    if "CSB Tournummer" in result.columns:
        result["__tour_sort"] = sort_key(result["CSB Tournummer"])
        sort_columns = ["__tour_sort", "CSB Tournummer"]
        if "CSB Kundennummer" in result.columns:
            result["__kunde_sort"] = sort_key(result["CSB Kundennummer"])
            sort_columns.extend(["__kunde_sort", "CSB Kundennummer"])
        result = result.sort_values(sort_columns, kind="mergesort")
        result = result.drop(columns=[column for column in ["__tour_sort", "__kunde_sort"] if column in result.columns])
    return result.reset_index(drop=True)


def calculate_unit_row(row: pd.Series) -> int:
    e2 = row["Anzahl gepl. E2"]
    e1 = row["Anzahl gepl. E1"]
    karton = row["Anzahl gepl. KARTON"]

    if row["Einheit"] == "TKT":
        # TKT: alles voll addieren und durch 12 teilen.
        return round_up_full((e2 + e1 + karton) / 12)

    # Rolli: E2 voll, E1 halb, KARTON halb, danach durch 16 teilen.
    return round_up_full((e2 + e1 * 0.5 + karton * 0.5) / 16)


def prepare_data(df: pd.DataFrame):
    found_columns, missing_columns = find_required_columns(df)
    if missing_columns:
        return None, None, missing_columns, found_columns

    work = pd.DataFrame({target: df[source] for target, source in found_columns.items()})

    for column in ["CSB Tournummer", "Kundenname", "Stadt", "Straße"]:
        work[column] = work[column].apply(clean_text)

    work["CSB Kundennummer"] = work["CSB Kundennummer"].apply(clean_customer_number)

    for column in NUMBER_COLUMNS:
        work[column] = work[column].apply(parse_number)

    grouped = (
        work.groupby(
            ["CSB Tournummer", "CSB Kundennummer", "Kundenname", "Stadt", "Straße"],
            dropna=False,
            as_index=False,
        )[NUMBER_COLUMNS]
        .sum()
    )

    for column in NUMBER_COLUMNS:
        grouped[column] = grouped[column].round(0).astype(int)

    grouped["Einheit"] = grouped["CSB Tournummer"].apply(lambda value: "TKT" if is_tkt_tour(value) else "Rolli")
    grouped[CALCULATED_COLUMN] = grouped.apply(calculate_unit_row, axis=1).astype(int)

    overview = (
        grouped.groupby(["CSB Tournummer", "Einheit"], dropna=False, as_index=False)[NUMBER_COLUMNS + [CALCULATED_COLUMN]]
        .sum()
    )

    customer_counts = (
        grouped.groupby(["CSB Tournummer", "Einheit"], dropna=False)["CSB Kundennummer"]
        .nunique()
        .reset_index(name="Anzahl Kunden")
    )

    overview = overview.merge(customer_counts, on=["CSB Tournummer", "Einheit"], how="left")
    overview = overview[["CSB Tournummer", "Einheit", "Anzahl Kunden"] + NUMBER_COLUMNS + [CALCULATED_COLUMN]]

    grouped = sort_result(grouped)[OUTPUT_COLUMNS]
    overview = sort_result(overview)

    for column in ["Anzahl Kunden"] + NUMBER_COLUMNS + [CALCULATED_COLUMN]:
        if column in overview.columns:
            overview[column] = overview[column].fillna(0).astype(int)

    return grouped, overview, [], found_columns


# -----------------------------------------------------------------------------
# Excel-Erzeugung
# -----------------------------------------------------------------------------
def autosize_worksheet(worksheet) -> None:
    for column_cells in worksheet.columns:
        column_letter = get_column_letter(column_cells[0].column)
        max_length = 0
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, len(value))
        worksheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 42)


def style_worksheet(worksheet) -> None:
    """
    Filter wird bewusst als normaler AutoFilter gesetzt.
    Keine Excel-Tabellenobjekte, damit Excel keine Reparaturmeldung erzeugt.
    """
    header_fill = PatternFill("solid", fgColor="1F2937")
    header_font = Font(color="FFFFFF", bold=True)
    even_fill = PatternFill("solid", fgColor="F9FAFB")
    tkt_fill = PatternFill("solid", fgColor="FEF3C7")
    thin_gray = Side(style="thin", color="D1D5DB")
    border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)

    worksheet.freeze_panes = "A2"

    if worksheet.max_row >= 1 and worksheet.max_column >= 1:
        last_column = get_column_letter(worksheet.max_column)
        worksheet.auto_filter.ref = f"A1:{last_column}{worksheet.max_row}"

    headers = [cell.value for cell in worksheet[1]]
    einheit_index = headers.index("Einheit") + 1 if "Einheit" in headers else None

    for row in worksheet.iter_rows():
        is_tkt_row = False
        if einheit_index and row[0].row > 1:
            is_tkt_row = worksheet.cell(row=row[0].row, column=einheit_index).value == "TKT"

        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=False)
            if cell.row == 1:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            elif is_tkt_row:
                cell.fill = tkt_fill
            elif cell.row % 2 == 0:
                cell.fill = even_fill

    worksheet.row_dimensions[1].height = 30

    number_headers = set(NUMBER_COLUMNS + [CALCULATED_COLUMN, "Anzahl Kunden"])
    for header_cell in worksheet[1]:
        if header_cell.value in number_headers:
            column_letter = get_column_letter(header_cell.column)
            for value_cell in worksheet[column_letter][1:]:
                value_cell.number_format = "0"
                value_cell.alignment = Alignment(horizontal="right", vertical="top")

    autosize_worksheet(worksheet)


def build_excel(grouped: pd.DataFrame, overview: pd.DataFrame) -> bytes:
    output = io.BytesIO()

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        grouped.to_excel(writer, sheet_name="Kunden je CSB Tour", index=False)
        overview.to_excel(writer, sheet_name="Übersicht je CSB Tour", index=False)

        for worksheet in writer.book.worksheets:
            style_worksheet(worksheet)

    output.seek(0)
    return output.getvalue()


def format_for_screen(df: pd.DataFrame) -> pd.DataFrame:
    result = df.copy()
    for column in NUMBER_COLUMNS + [CALCULATED_COLUMN, "Anzahl Kunden"]:
        if column in result.columns:
            result[column] = result[column].fillna(0).astype(int)
    return result


# -----------------------------------------------------------------------------
# Oberfläche
# -----------------------------------------------------------------------------
st.title("Auftragspool Auswertung")
st.caption("CSV hochladen, nach CSB Tournummer und Kunde zusammenfassen, Rolli oder TKT berechnen und Excel herunterladen.")

uploaded_file = st.file_uploader(
    "Auftragspool CSV hochladen",
    type=["csv"],
    help="Erwartet wird der Kisoft Auftragspool mit Semikolon als Trennzeichen.",
)

if not uploaded_file:
    st.info("Bitte eine Auftragspool CSV hochladen.")
    st.stop()

try:
    raw_df = read_csv_robust(uploaded_file)
except Exception as error:
    st.error(str(error))
    st.stop()

grouped_df, overview_df, missing_columns, found_columns = prepare_data(raw_df)

if missing_columns:
    st.error("Die Datei enthält nicht alle benötigten Spalten.")
    st.write("Fehlende Spalten:")
    st.write(missing_columns)
    with st.expander("Gefundene Spalten anzeigen"):
        st.write(list(raw_df.columns))
    st.stop()

excel_bytes = build_excel(grouped_df, overview_df)
file_date = datetime.now().strftime("%Y_%m_%d")

rolli_sum = int(grouped_df.loc[grouped_df["Einheit"] == "Rolli", CALCULATED_COLUMN].sum())
tkt_sum = int(grouped_df.loc[grouped_df["Einheit"] == "TKT", CALCULATED_COLUMN].sum())

col1, col2, col3, col4, col5 = st.columns(5)
col1.metric("Zeilen in CSV", f"{len(raw_df):,}".replace(",", "."))
col2.metric("Kunden", f"{grouped_df['CSB Kundennummer'].nunique():,}".replace(",", "."))
col3.metric("CSB Tournummern", f"{grouped_df['CSB Tournummer'].nunique():,}".replace(",", "."))
col4.metric("Rollis gesamt", f"{rolli_sum:,}".replace(",", "."))
col5.metric("TKT gesamt", f"{tkt_sum:,}".replace(",", "."))

st.download_button(
    "Excel herunterladen",
    data=excel_bytes,
    file_name=f"Auftragspool_CSB_Tournummer_Kunden_Mengen_Rolli_TKT_{file_date}.xlsx",
    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    use_container_width=True,
)

st.divider()

left, right = st.columns([2, 1])
with left:
    selected_tours = st.multiselect(
        "CSB Tournummer filtern",
        options=sorted(grouped_df["CSB Tournummer"].dropna().unique(), key=lambda x: (not str(x).isdigit(), str(x))),
    )
with right:
    selected_units = st.multiselect(
        "Einheit filtern",
        options=["Rolli", "TKT"],
        default=[],
    )

view_df = grouped_df.copy()
if selected_tours:
    view_df = view_df[view_df["CSB Tournummer"].isin(selected_tours)]
if selected_units:
    view_df = view_df[view_df["Einheit"].isin(selected_units)]

st.subheader("Kunden je CSB Tour")
st.dataframe(format_for_screen(view_df), use_container_width=True, hide_index=True)

with st.expander("Übersicht je CSB Tour anzeigen", expanded=False):
    st.dataframe(format_for_screen(overview_df), use_container_width=True, hide_index=True)

with st.expander("TKT-Regel anzeigen", expanded=False):
    st.write("Exakte TKT-Touren:")
    st.write(sorted(TKT_EXACT_TOURS))
    st.write("Zusätzlich TKT, wenn die Tournummer fünfstellig ist und nach der ersten Ziffer eines dieser Muster hat:")
    st.write(sorted(TKT_SUFFIXES))

with st.expander("Erkannte CSV-Spalten anzeigen", expanded=False):
    st.write(found_columns)

st.markdown(
    "<div class='small-note'>Berechnung Rolli: "
    "(Anzahl gepl. E2 + 0,5 × Anzahl gepl. E1 + 0,5 × Anzahl gepl. KARTON) / 16 — danach auf volle Rollis aufgerundet. "
    "Berechnung TKT: (Anzahl gepl. E2 + Anzahl gepl. E1 + Anzahl gepl. KARTON) / 12 — danach auf volle TKT aufgerundet. "
    "TKT-Touren werden über die exakten Tournummern und den gleichen Tournummer-Aufbau erkannt. "
    "Die Excel-Datei nutzt normale Filter im Kopfbereich und keine Excel-Tabellenobjekte. Einzelaufträge werden nicht exportiert.</div>",
    unsafe_allow_html=True,
)
